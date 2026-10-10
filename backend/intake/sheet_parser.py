"""Layout-aware Excel parser synthesis and validation."""
from __future__ import annotations
import hashlib, json, time
from pathlib import Path
from pydantic import BaseModel
from backend import config
from backend.intake import sandbox
from backend.llm.client import call_json
from backend.llm.prompts import gen_sheet_parser
from backend.models import Evidence, ParserOut, Rule, Roster
from backend.registry import db

class ParseError(RuntimeError): pass
class IngestSheetResult(BaseModel):
    rules: list[Rule]
    cache_hit: bool
    attempts: int
    tokens_used: int
    seconds: float

def layout_signature(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc: raise ParseError("Install openpyxl to read spreadsheets") from exc
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    try:
        layout=[]
        for sheet in book.worksheets:
            headers=[]
            for row in sheet.iter_rows(min_row=1,max_row=10,values_only=True):
                vals=[str(v).strip().lower() for v in row if v is not None and str(v).strip()]
                if vals: headers=vals; break
            layout.append([sheet.title,headers])
    finally: book.close()
    return hashlib.sha256(json.dumps(layout,ensure_ascii=False,separators=(",",":")).encode()).hexdigest()[:16]

def _sample(path: Path) -> str:
    import openpyxl
    book=openpyxl.load_workbook(path,read_only=True,data_only=True); lines=[]
    try:
        sheet=book.active
        for row_no,row in enumerate(sheet.iter_rows(min_row=1,max_row=10,values_only=True),1):
            lines.append(f"{row_no}: " + " | ".join("" if v is None else str(v) for v in row))
    finally: book.close()
    return "\n".join(lines)

def _validated(rows: list[dict], roster: Roster) -> tuple[list[tuple[str,list[str],str]],int]:
    course_map={s.course.strip().casefold():s for s in roster.sessions}; teachers=set(roster.teachers); good=[]; failures=0
    for row in rows:
        try:
            if not isinstance(row.get("course"),str) or not isinstance(row.get("teachers"),list) or not isinstance(row.get("source_cell"),str): raise ValueError
            session=course_map[row["course"].strip().casefold()]
            names=row["teachers"]
            if not names or any(not isinstance(t,str) or t not in teachers or t not in session.teachers for t in names): raise ValueError
            good.append((session.id,list(dict.fromkeys(names)),row["source_cell"]))
        except (KeyError,ValueError,TypeError): failures+=1
    return good,failures

def ingest_sheet(path: Path, filename: str, roster: Roster | None=None) -> IngestSheetResult:
    start=time.monotonic(); tokens_start=time.time(); path=Path(path)
    if path.stat().st_size>50*1024*1024: raise ParseError("Spreadsheet exceeds 50 MiB")
    roster=roster or db.get_roster(); sig=layout_signature(path); code=db.get_parser(sig); cache_hit=code is not None; attempts=0; last_error=""
    if code:
        try:
            rows=sandbox.run_parser(code,path); good,failed=_validated(rows,roster)
            if not good or failed/len(rows)>.05: raise ParseError("cached parser output did not validate")
        except Exception as exc:
            db.delete_parser(sig); code=None; cache_hit=False; last_error=str(exc)
    if code is None:
        for attempt in range(1,4):
            attempts=attempt
            generated=call_json("gen_sheet_parser","reason",[{"role":"system","content":"Return a standalone Python parser as JSON."},{"role":"user","content":gen_sheet_parser(_sample(path),last_error)}],ParserOut)
            code=generated.code
            try:
                rows=sandbox.run_parser(code,path); good,failed=_validated(rows,roster)
                if not rows or not good or failed/len(rows)>.05: raise ParseError(f"{failed} of {len(rows)} rows failed validation")
                db.save_parser(sig,code); break
            except Exception as exc:
                last_error=str(exc); code=None
        if code is None: raise ParseError(f"Could not parse spreadsheet after {attempts} attempts: {last_error}")
    if not rows or not good or failed/len(rows)>.05: raise ParseError(f"{failed} of {len(rows)} rows failed validation")
    rules=[]
    for sid,names,cell in good:
        rule=Rule(id=db.next_rule_id(),type="only_qualified",params={"session_id":sid,"teachers":names},owner="Dean",evidence=[Evidence(kind="sheet",ref=f"{Path(filename).name}!{cell}")])
        db.save_rule(rule); rules.append(rule)
    return IngestSheetResult(rules=rules,cache_hit=cache_hit,attempts=attempts,tokens_used=db.tokens_since(tokens_start),seconds=time.monotonic()-start)
